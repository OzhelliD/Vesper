"""
attachments.py — Convierte archivos adjuntos en bloques que Claude entiende.

Soporta:
  - Imágenes (jpg, png, gif, webp)         → bloque "image"
  - PDF                                    → bloque "document" (Claude lee texto y gráficas)
  - Word (.docx)                           → texto extraído (párrafos + tablas)
  - Excel (.xlsx, .xlsm)                   → texto tipo CSV por hoja (recortado)
  - Texto plano (.txt, .csv, .md, .json)   → texto

El frontend manda cada archivo como:
  {"name": "ticket.pdf", "media_type": "application/pdf", "data": "<base64>"}
"""

import base64, csv, io, os

MAX_FILES          = 5
MAX_FILE_BYTES     = 20 * 1024 * 1024   # 20 MB por archivo (antes de base64)
MAX_IMAGE_BYTES    = 5 * 1024 * 1024    # límite de la API por imagen
MAX_TEXT_CHARS     = 60_000             # por archivo de texto extraído
MAX_EXCEL_ROWS     = 400                # por hoja
MAX_EXCEL_SHEETS   = 8

IMAGE_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}
PDF_TYPES   = {"application/pdf"}
DOCX_TYPES  = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
XLSX_TYPES  = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
               "application/vnd.ms-excel.sheet.macroenabled.12"}
TEXT_EXT    = {".txt", ".csv", ".md", ".json", ".log", ".tsv"}

EXT_TO_TYPE = {
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
    ".gif": "image/gif", ".webp": "image/webp", ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroenabled.12",
}


class AttachmentError(ValueError):
    """Error legible para el usuario."""


def _guess_type(name, declared):
    ext = os.path.splitext(name or "")[1].lower()
    declared = (declared or "").lower().strip()
    if declared and declared != "application/octet-stream":
        if declared == "image/jpg":
            return "image/jpeg"
        return declared
    return EXT_TO_TYPE.get(ext, "text/plain" if ext in TEXT_EXT else declared)


def _clip(text, limit=MAX_TEXT_CHARS):
    text = (text or "").strip()
    if len(text) > limit:
        return text[:limit] + f"\n\n[... recortado, el archivo tenía {len(text):,} caracteres ...]"
    return text


def _docx_to_text(raw):
    from docx import Document  # python-docx
    doc = Document(io.BytesIO(raw))
    out = []
    for p in doc.paragraphs:
        if p.text.strip():
            style = (p.style.name or "").lower() if p.style is not None else ""
            prefix = "# " if style.startswith("heading") or style.startswith("título") else ""
            out.append(prefix + p.text.strip())
    for ti, table in enumerate(doc.tables, 1):
        out.append(f"\n[Tabla {ti}]")
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _xlsx_to_text(raw):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out = []
    for si, ws in enumerate(wb.worksheets):
        if si >= MAX_EXCEL_SHEETS:
            out.append(f"\n[... {len(wb.worksheets) - MAX_EXCEL_SHEETS} hojas más omitidas ...]")
            break
        out.append(f"\n## Hoja: {ws.title}")
        buf = io.StringIO()
        writer = csv.writer(buf)
        rows_written = 0
        total_rows = 0
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v is None or str(v).strip() == "" for v in row):
                continue
            total_rows += 1
            if rows_written < MAX_EXCEL_ROWS:
                writer.writerow(["" if v is None else v for v in row])
                rows_written += 1
        out.append(buf.getvalue().rstrip())
        if total_rows > rows_written:
            out.append(f"[... {total_rows - rows_written} filas más omitidas ...]")
    wb.close()
    return "\n".join(out)


def _decode_text(raw):
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def build_attachment_blocks(attachments):
    """
    Recibe la lista del frontend y regresa (blocks, labels, errors):
      blocks → lista de bloques de contenido para la API de Claude
      labels → nombres para guardar en el historial ("ticket.pdf")
      errors → avisos legibles de archivos que no se pudieron leer
    """
    blocks, labels, errors = [], [], []
    for item in (attachments or [])[:MAX_FILES]:
        name = (item or {}).get("name") or "archivo"
        data = (item or {}).get("data") or ""
        mtype = _guess_type(name, item.get("media_type"))
        try:
            raw = base64.b64decode(data, validate=False)
        except Exception:
            errors.append(f"{name}: no pude decodificar el archivo.")
            continue
        if not raw:
            errors.append(f"{name}: el archivo está vacío.")
            continue
        if len(raw) > MAX_FILE_BYTES:
            errors.append(f"{name}: pesa más de {MAX_FILE_BYTES // (1024*1024)} MB.")
            continue

        try:
            if mtype in IMAGE_TYPES:
                if len(raw) > MAX_IMAGE_BYTES:
                    errors.append(f"{name}: la imagen pesa más de 5 MB.")
                    continue
                blocks.append({"type": "image",
                               "source": {"type": "base64", "media_type": mtype, "data": data}})
            elif mtype in PDF_TYPES:
                blocks.append({"type": "document",
                               "source": {"type": "base64", "media_type": "application/pdf", "data": data},
                               "title": name})
            elif mtype in DOCX_TYPES or name.lower().endswith(".docx"):
                text = _clip(_docx_to_text(raw))
                blocks.append({"type": "text", "text": f"[Archivo Word: {name}]\n{text or '(sin texto)'}"})
            elif mtype in XLSX_TYPES or name.lower().endswith((".xlsx", ".xlsm")):
                text = _clip(_xlsx_to_text(raw))
                blocks.append({"type": "text", "text": f"[Archivo Excel: {name}]\n{text or '(sin datos)'}"})
            elif mtype.startswith("text/") or mtype == "application/json" \
                    or os.path.splitext(name)[1].lower() in TEXT_EXT:
                text = _clip(_decode_text(raw))
                blocks.append({"type": "text", "text": f"[Archivo: {name}]\n{text}"})
            elif name.lower().endswith((".doc", ".xls")):
                errors.append(f"{name}: formato antiguo de Office; guárdelo como .docx o .xlsx.")
                continue
            elif mtype in ("image/heic", "image/heif"):
                errors.append(f"{name}: formato HEIC no soportado; envíelo como JPG.")
                continue
            else:
                errors.append(f"{name}: tipo de archivo no soportado ({mtype or 'desconocido'}).")
                continue
            labels.append(name)
        except Exception as e:
            print(f"[ATTACH] Error leyendo {name}: {e}")
            errors.append(f"{name}: no pude leer el contenido.")
    if attachments and len(attachments) > MAX_FILES:
        errors.append(f"Solo se procesan {MAX_FILES} archivos por mensaje.")
    return blocks, labels, errors
